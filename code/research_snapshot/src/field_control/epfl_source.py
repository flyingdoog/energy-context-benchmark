"""Read official EPFL workbooks into the cache consumed by the task builder."""
import json
from pathlib import Path
import time
from urllib.request import Request, urlopen
import pandas as pd

REVISION = '680bc3bb751f86a47db32a5641f09dcfa079caae'
SOURCE_URL = 'https://github.com/DESL-EPFL/Level-3-EV-charging-dataset'
RAW_BASE = f'https://raw.githubusercontent.com/DESL-EPFL/Level-3-EV-charging-dataset/{REVISION}/'
FILES = ('README.md', 'LICENSE', 'Measurement_data.xlsx', 'Session_data.xlsx')


def download_source(destination):
    """Download a fixed public release locally; retain its original notice."""
    dest = Path(destination).resolve()
    dest.mkdir(parents=True, exist_ok=True)
    record_path = dest / 'acquisition.json'
    record = json.loads(record_path.read_text()) if record_path.exists() else {
        'source_url': SOURCE_URL, 'revision': REVISION, 'files': []}
    if record.get('revision') != REVISION or record.get('source_url') != SOURCE_URL:
        raise ValueError('Source receipt belongs to another release')
    known = {r['name']: r for r in record['files']}
    for name in FILES:
        target = dest / name
        if target.exists():
            if name not in known or target.stat().st_size != known[name]['bytes']:
                raise ValueError(f'Existing source file needs inspection: {target}')
            continue
        url = RAW_BASE + name
        temporary = target.with_name(target.name + '.partial')
        total = 0
        with urlopen(Request(url, headers={'User-Agent': 'OperationalForecastBenchmark/0.1'}), timeout=30) as response:
            with temporary.open('wb') as out:
                while block := response.read(1024 * 1024):
                    total += len(block)
                    if total > 32 * 1024 * 1024:
                        raise ValueError('Source file exceeds the 32 MiB acquisition bound')
                    out.write(block)
        if not total:
            raise ValueError(f'Empty source download: {name}')
        temporary.rename(target)
        known[name] = {'name': name, 'bytes': total, 'url': url}
        record.update(files=list(known.values()), completed=len(known) == len(FILES))
        record_path.write_text(json.dumps(record, indent=2) + '\n', encoding='utf-8')
        print(json.dumps({'downloaded': name, 'bytes': total}), flush=True)
    return record


def read_measurements(path):
    """One streaming pass: ignore empty calendar slots, preserve observed zeros."""
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    parts = []
    support = []
    try:
        for name in ('CCS1', 'CCS2'):
            iterator = book[name].iter_rows(values_only=True)
            headers = list(next(iterator))
            if 'Session' not in headers:
                raise ValueError(f'Missing Session column in {name}')
            session = headers.index('Session')
            active = []
            total = 0
            for row in iterator:
                total += 1
                if row[session] is not None:
                    active.append(row)
            parts.append(pd.DataFrame(active, columns=headers))
            support.append({'sheet': name, 'calendar_rows': total, 'active_rows': len(active)})
            print(json.dumps(support[-1]), flush=True)
    finally:
        book.close()
    return pd.concat(parts, ignore_index=True), support


def link_sessions(data, sessions):
    """Link plug-local measurement IDs to global session IDs by timestamps."""
    data = data.copy()
    sessions = sessions.copy()
    for frame in (data, sessions):
        frame['key'] = frame.CCS.astype(str) + ':' + frame.Session.astype(int).astype(str)
    data['Date and time'] = pd.to_datetime(data['Date and time'])
    sessions['Arrival'] = pd.to_datetime(sessions.Arrival)
    sessions['Departure'] = pd.to_datetime(sessions.Departure)
    if data.duplicated(['key', 'Date and time']).any():
        raise ValueError('Duplicate within-session timestamps')
    if not sessions.key.is_unique:
        raise ValueError('Duplicate source session keys')
    starts = data.groupby(['CCS', 'Session', 'key'])['Date and time'].agg(['min', 'max']).reset_index()
    linked = starts.merge(sessions[['CCS', 'Session', 'key', 'Arrival', 'Departure']],
        left_on=['CCS', 'min'], right_on=['CCS', 'Arrival'], how='outer',
        suffixes=('_measurement', '_session'), validate='one_to_one', indicator=True)
    if not (linked['_merge'] == 'both').all() or not (linked['max'] == linked.Departure).all():
        raise ValueError('Source session endpoints do not establish one-to-one linkage')
    support = pd.DataFrame({
        'key': linked.key_measurement,
        'start_offset_seconds': (linked['min'] - linked.Arrival).dt.total_seconds(),
        'end_offset_seconds': (linked['max'] - linked.Departure).dt.total_seconds()})
    return data.drop(columns='key'), support, linked


def extract_source(source, destination):
    """Build only the source cache required for query preparation, without models."""
    source = Path(source).resolve()
    dest = Path(destination).resolve()
    if source == dest:
        raise ValueError('Keep source workbooks and prepared cache in separate folders')
    receipt_path = dest / 'source_receipt.json'
    outputs = ['source_active_rows.csv.gz', 'session_support.csv', 'session_id_crosswalk.csv']
    if receipt_path.exists():
        record = json.loads(receipt_path.read_text())
        if record.get('source_directory') != str(source) or not record.get('completed'):
            raise ValueError('Cache receipt belongs to another or incomplete source')
        if not all((dest / name).exists() for name in outputs):
            raise ValueError('Completed cache has missing files')
        return record
    if dest.exists() and any(dest.iterdir()):
        raise ValueError('Incomplete extraction directory needs inspection; choose a new work directory')
    for name in ('Measurement_data.xlsx', 'Session_data.xlsx'):
        if not (source / name).is_file():
            raise FileNotFoundError(source / name)
    started = time.monotonic()
    data, sheets = read_measurements(source / 'Measurement_data.xlsx')
    sessions = pd.read_excel(source / 'Session_data.xlsx', sheet_name='Session_data')
    data, support, crosswalk = link_sessions(data, sessions)
    dest.mkdir(parents=True, exist_ok=True)
    data.to_csv(dest / outputs[0], index=False)
    support.to_csv(dest / outputs[1], index=False)
    crosswalk.to_csv(dest / outputs[2], index=False)
    record = dict(completed=True, source_directory=str(source), source_url=SOURCE_URL,
        sheets=sheets, active_rows=len(data), sessions=len(support),
        empty_calendar_rows=sum(r['calendar_rows'] - r['active_rows'] for r in sheets),
        endpoint_linkage='One-to-one by plug/arrival and confirmed departure',
        source_modified=False, new_model_fits=0, seconds=time.monotonic() - started)
    receipt_path.write_text(json.dumps(record, indent=2) + '\n', encoding='utf-8')
    return record
